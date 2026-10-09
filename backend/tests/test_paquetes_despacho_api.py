"""Tests de integración de /api/paquetes-despacho (Paso 1.1): cargar el Excel de despacho,
listar, corregir direcciones sin sector y exportar. Necesitan la BD con las migraciones
031–033. El permiso de página se reemplaza por un usuario fijo para no depender de los
usuarios sembrados en la BD; aparte se prueba que sin token se rechaza."""
import io
from datetime import date

import openpyxl
import pytest
from sqlalchemy import delete

from app.database import AsyncSessionLocal
from app.main import app
from app.models.paquetes_despacho import PaqueteDespacho
from app.routers import paquetes_despacho as router_module

PREFIJO = "TEST11-"
F_EMI = date(2026, 10, 8)
URL = "/api/paquetes-despacho"
COLUMNAS_EN = ["Waybill number", "Recipient's name", "Customer phone", "Address2", "Otra columna"]
COLUMNAS_ES = ["Número de Guía", "El nombre del destinatario", "Teléfono entrante",
               "Dirección detallada del destinatario"]


def _excel(columnas: list[str], filas: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(columnas)
    for fila in filas:
        ws.append(fila)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _archivo(contenido: bytes, nombre: str = "despacho.xlsx"):
    return {"file": (nombre, contenido, "application/octet-stream")}


@pytest.fixture(autouse=True)
async def limpiar_y_autorizar():
    usuario = {"username": "test", "rol": "administrador"}
    app.dependency_overrides[router_module._auth.dependency] = lambda: usuario
    app.dependency_overrides[router_module._auth_destino.dependency] = lambda: usuario
    await _limpiar()
    yield
    await _limpiar()
    app.dependency_overrides.pop(router_module._auth.dependency, None)
    app.dependency_overrides.pop(router_module._auth_destino.dependency, None)


async def _limpiar():
    async with AsyncSessionLocal() as db:
        await db.execute(delete(PaqueteDespacho).where(PaqueteDespacho.serial.like(f"{PREFIJO}%")))
        await db.commit()


async def _cargar(client, contenido: bytes, f_emi: date = F_EMI, nombre: str = "despacho.xlsx"):
    return await client.post(f"{URL}/cargar", files=_archivo(contenido, nombre), data={"f_emi": f_emi.isoformat()})


# ── Cargar ────────────────────────────────────────────────────────────────────

async def test_cargar_excel_en_ingles(client):
    contenido = _excel(COLUMNAS_EN, [
        [f"{PREFIJO}1", "Ana", "3001112233", "calle 95 # 49-22 clinica del pie y spa", "x"],
        [f"{PREFIJO}2", "Luis", "3004445566", "CR 20 # 66-15", "y"],
        [f"{PREFIJO}3", "Eva", None, "direccion rara", None],
    ])
    r = await _cargar(client, contenido)

    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["total"], body["creados"], body["reemplazados"]) == (3, 3, 0)
    assert body["sin_sector"] == [{"serial": f"{PREFIJO}3", "direccion": "direccion rara"}]
    assert body["devoluciones"] == 1  # el que no cae en ninguna zona
    assert body["f_emi"] == F_EMI.isoformat()


async def test_cargar_excel_en_espanol_y_volver_a_cargar_reemplaza(client):
    contenido = _excel(COLUMNAS_ES, [[f"{PREFIJO}1", "Ana", "300", "CL 82 38 10"]])
    assert (await _cargar(client, contenido)).json()["creados"] == 1

    r = await _cargar(client, contenido, f_emi=date(2026, 10, 9))
    assert (r.json()["creados"], r.json()["reemplazados"]) == (0, 1)


async def test_cargar_base_de_whatsapp_con_columna_address(client):
    # Así llega la base de despacho por WhatsApp (DESPACHO SERVILLA 0810.xlsx): "Address", sin "Address2"
    columnas = ["Waybill number", "Destino", "Recipient's name", "Address", "Customer phone"]
    contenido = _excel(columnas, [[f"{PREFIJO}1", "BOG-Doce de Octubre.DS", "Ana", "carrera 27 #76-54 Bodega", "573202511797"]])
    r = await _cargar(client, contenido)

    assert r.status_code == 200, r.text
    assert (r.json()["creados"], r.json()["sin_sector"]) == (1, [])
    paquetes = (await client.get(f"{URL}/", params={"f_emi": F_EMI.isoformat()})).json()
    assert next(p for p in paquetes if p["serial"] == f"{PREFIJO}1")["direccion"] == "carrera 27 #76-54 Bodega"


async def test_con_address2_y_address_manda_address2(client):
    columnas = ["Waybill number", "Recipient's name", "Customer phone", "Address", "Address2"]
    await _cargar(client, _excel(columnas, [[f"{PREFIJO}1", "Ana", "300", "texto viejo", "CL 82 38 10"]]))

    paquetes = (await client.get(f"{URL}/", params={"f_emi": F_EMI.isoformat()})).json()
    assert next(p for p in paquetes if p["serial"] == f"{PREFIJO}1")["direccion"] == "CL 82 38 10"


async def test_serial_numerico_no_se_vuelve_decimal(client):
    contenido = _excel(COLUMNAS_EN, [[9_900_000_000_001, "Ana", 3001112233, "CL 82 38 10", None]])
    await _cargar(client, contenido)

    r = await client.get(f"{URL}/", params={"f_emi": F_EMI.isoformat(), "zona": "80_3"})
    seriales = [p["serial"] for p in r.json()]
    assert "9900000000001" in seriales
    async with AsyncSessionLocal() as db:
        await db.execute(delete(PaqueteDespacho).where(PaqueteDespacho.serial == "9900000000001"))
        await db.commit()


async def test_columnas_faltantes_dice_cuales_faltan_y_cuales_trae(client):
    r = await _cargar(client, _excel(["Waybill number", "Address2", "Ciudad"], [[f"{PREFIJO}1", "CL 82 38 10", "Bogota"]]))

    assert r.status_code == 400
    detalle = r.json()["detail"]
    assert detalle["faltantes"] == [
        "nombre (esperado: Recipient's name o El nombre del destinatario)",
        "telefono (esperado: Customer phone o Teléfono entrante)",
    ]
    assert detalle["columnas_archivo"] == ["Waybill number", "Address2", "Ciudad"]


@pytest.mark.parametrize("nombre,contenido,mensaje", [
    ("despacho.csv", b"serial\n1", "Solo se aceptan archivos Excel"),
    ("despacho.xlsx", b"esto no es un excel", "No se pudo leer el archivo Excel"),
])
async def test_archivo_invalido(client, nombre, contenido, mensaje):
    r = await _cargar(client, contenido, nombre=nombre)
    assert r.status_code == 400
    assert mensaje in r.json()["detail"]


async def test_excel_vacio(client):
    r = await _cargar(client, _excel(COLUMNAS_EN, []))
    assert r.status_code == 400
    assert "vacío" in r.json()["detail"]


# ── Listar, corregir y exportar ───────────────────────────────────────────────

async def _cargar_tres(client):
    await _cargar(client, _excel(COLUMNAS_EN, [
        [f"{PREFIJO}1", "Ana", "300", "calle 95 # 49-22", None],
        [f"{PREFIJO}2", "Luis", "301", "CR 20 # 66-15", None],
        [f"{PREFIJO}3", "Eva", "302", "Calle Ac11sur#16este99", None],
    ]))


async def test_listar_por_fecha_zona_y_sin_sector(client):
    await _cargar_tres(client)
    params = {"f_emi": F_EMI.isoformat()}

    todos = (await client.get(f"{URL}/", params=params)).json()
    assert {p["serial"] for p in todos} >= {f"{PREFIJO}1", f"{PREFIJO}2", f"{PREFIJO}3"}

    zona = (await client.get(f"{URL}/", params={**params, "zona": "60_4"})).json()
    assert [p["serial"] for p in zona if p["serial"].startswith(PREFIJO)] == [f"{PREFIJO}2"]

    sin_sector = (await client.get(f"{URL}/", params={**params, "solo_sin_sector": True})).json()
    assert [p["serial"] for p in sin_sector if p["serial"].startswith(PREFIJO)] == [f"{PREFIJO}3"]


async def test_corregir_direccion_vuelve_a_sectorizar(client):
    await _cargar_tres(client)

    r = await client.patch(f"{URL}/{PREFIJO}3/direccion", json={"direccion": "CR 17 # 63-49"})

    assert r.status_code == 200, r.text
    p = r.json()
    assert (p["direccion"], p["direccion_estandarizada"]) == ("CR 17 # 63-49", "CR 17 63 49")
    assert (p["localidad"], p["zona"]) == ("Chapinero", "60_2")


async def test_corregir_serial_inexistente(client):
    r = await client.patch(f"{URL}/{PREFIJO}NO-EXISTE/direccion", json={"direccion": "CL 82 38 10"})
    assert r.status_code == 404


@pytest.mark.parametrize("formato,tipo", [
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("csv", "text/csv"),
])
async def test_exportar(client, formato, tipo):
    await _cargar_tres(client)

    r = await client.get(f"{URL}/exportar", params={"f_emi": F_EMI.isoformat(), "formato": formato})

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(tipo)
    assert f"paquetes_sectorizados_{F_EMI.isoformat()}.{formato}" in r.headers["content-disposition"]
    if formato == "csv":
        texto = r.content.decode("utf-8-sig")
        assert texto.splitlines()[0] == "serial,nombre,telefono,direccion,direccion_estandarizada,codigo_postal,localidad,zona,f_emi,estado"
        assert f"{PREFIJO}2,Luis,301,CR 20 # 66-15,CR 20 66 15,110231,Chapinero,60_4,2026-10-08,sin gestión" in texto
    else:
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        seriales = [ws.cell(row=i, column=1).value for i in range(4, ws.max_row + 1)]
        assert f"{PREFIJO}1" in seriales


async def test_devoluciones_son_los_fuera_de_zona_con_el_formato_de_devoluciones_enriquecidas(client):
    await _cargar_tres(client)

    r = await client.get(f"{URL}/devoluciones", params={"f_emi": F_EMI.isoformat()})

    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f"devoluciones_{F_EMI.isoformat()}.xlsx" in r.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    filas = [fila for fila in ws.iter_rows(values_only=True) if fila[0] == "serial" or str(fila[0]).startswith(PREFIJO)]
    assert filas == [
        ("serial", "nombre", "telefono", "direccion", "localidad"),
        (f"{PREFIJO}3", "Eva", "302", "Calle Ac11sur#16este99", None),  # fuera de zona (y sin localidad)
    ]
    assert ws["A1"].font.b


# ── Destino al escanear (Paso 2.6) ────────────────────────────────────────────

async def test_destino_de_un_paquete_con_zona(client):
    await _cargar_tres(client)

    r = await client.get(f"{URL}/destino/ {PREFIJO}2 ")  # el lector puede mandar espacios

    assert r.status_code == 200
    assert r.json() == {
        "serial": f"{PREFIJO}2", "ultimos_4": "11-2",
        "en_tabla": True, "direccion": "CR 20 # 66-15", "direccion_estandarizada": "CR 20 66 15",
        "localidad": "Chapinero",
        "zona": "60_4", "fuera_de_zona": False, "aviso": None,
    }


async def test_destino_fuera_de_zona(client):
    await _cargar(client, _excel(COLUMNAS_EN, [[f"{PREFIJO}9", "Ana", "300", "CL 21 33 40", None]]))

    body = (await client.get(f"{URL}/destino/{PREFIJO}9")).json()

    assert (body["localidad"], body["zona"]) == ("Puente Aranda", None)
    assert body["fuera_de_zona"] is True
    assert body["aviso"] == "Fuera de zona"


async def test_destino_de_un_serial_que_no_esta_en_la_tabla(client):
    body = (await client.get(f"{URL}/destino/{PREFIJO}NUEVO1234")).json()

    assert body["en_tabla"] is False
    assert body["ultimos_4"] == "1234"
    assert body["aviso"] == "No está en la tabla"
    assert (body["direccion"], body["zona"]) == (None, None)


# ── Permisos ──────────────────────────────────────────────────────────────────

async def test_sin_token_se_rechaza(client):
    app.dependency_overrides.pop(router_module._auth.dependency, None)
    r = await client.get(f"{URL}/", params={"f_emi": F_EMI.isoformat()})
    assert r.status_code in (401, 403)
