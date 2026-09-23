import io
import random

import pytest
from openpyxl import load_workbook

from app.services.formato_servilla_service import (
    causal_de_motivo,
    generar_formato_servilla,
    motivo_histo,
)
from app.services.generar_dat_service import generar_dat


@pytest.mark.parametrize(
    ("motivo", "causal"),
    [
        ("Entrega", "00"),
        ("Desconocido", "01"),
        ("Destinatario N.", "01"),
        ("Rehusada", "02"),
        ("REHUSADO", "02"),
        ("Traslado", "03"),
        ("Direccion Errad", "05"),
        ("Dirección Errada", "05"),
        (" Direccion Errada", "05"),
        ("No Cubrimiento", "05"),
        ("Direccion Incom", "06"),
        ("Dir Incompleta", "06"),
        ("Inconsistente", "07"),
        ("Zona Alto Riesg", "08"),
        ("Sobrante", "09"),
        ("Docto en Camino", "09"),
        ("Docto en Camino Pend", "09"),
        ("Cerrado", "13"),
        ("Siniestro", "14"),
        ("Destinatario De", ""),
        ("No Reparto", ""),
        ("Dev.Inicial", ""),
        ("Lleva Mensajero", ""),
        ("", ""),
    ],
)
def test_causal_de_motivo(motivo, causal):
    assert causal_de_motivo(motivo) == causal


def test_motivo_histo_regla_update_histo():
    assert motivo_histo({"retorno": "D", "ret_esc": "D", "motivo": "Traslado"}) == "Traslado"
    assert motivo_histo({"retorno": "l", "ret_esc": "E", "motivo": "Lleva Mensajero"}) == "Entrega"
    assert motivo_histo({"retorno": "o", "ret_esc": "", "motivo": ""}) == "Dirección Errada"
    assert motivo_histo({"retorno": "j", "ret_esc": "i", "cod_sec": "*DEV_22"}) == "Rehusada"
    assert motivo_histo({"retorno": "j", "ret_esc": "i", "motivo": "cajonera", "cod_sec": "_ARAUCA"}) == ""


def _fila(serial, courrier, retorno, motivo=""):
    return {"serial": serial, "courrier": courrier, "retorno": retorno, "ret_esc": "",
            "motivo": motivo, "cod_sec": "", "identdes": "1", "oficina": serial[-3:],
            "nombred": "X", "dir_pred": "CL 1", "barrd1": "", "ciudad1": "BOGOTA",
            "dpto1": "CUND", "orden": "123791"}


HISTO = [
    _fila("1672696003", "F&S", "E"),
    _fila("1672696010", "SERVILLA", "D", "Direccion Errad"),
    _fila("1672696027", "SERVILLA", "D", "No Reparto"),
    _fila("1672696034", "SERVILLA", "j"),
    _fila("1672693019", "PRINDEL", "E"),
    _fila("1672693026", " lecta ", "E"),
]


def test_formato_servilla_filtra_y_llena_columnas():
    r = generar_formato_servilla("123791", "20260803", HISTO, random.Random(1))
    ws = load_workbook(io.BytesIO(r.contenido)).active
    filas = [[c.value for c in fila] for fila in ws.iter_rows()]

    assert r.nombre == "formato_servilla_123791_20260803.xlsx"
    assert r.filas == 4 and r.excluidos == 2
    assert filas[0] == ["serial", "Estado", "Causal_Dev", "F_recepcio", "guias", "F_GESTION", "motivo"]
    por_serial = {f[0]: f for f in filas[1:]}
    assert set(por_serial) == {"1672696003", "1672696010", "1672696027", "1672696034"}
    assert por_serial["1672696003"][:5] == ["1672696003", "ENT", "00", "20260803", "1672696003"]
    assert por_serial["1672696010"][1:3] == ["DEV", "05"]
    assert por_serial["1672696027"][1:3] == [None, None]
    assert por_serial["1672696027"][6] == "No Reparto"
    for f in filas[1:]:
        assert "20260805" <= f[5] <= "20260809"
    assert [s["serial"] for s in r.sin_causal] == ["1672696027", "1672696034"]
    assert ws["C2"].number_format == "@"


def test_f_gestion_suma_entre_2_y_6_dias_calendario():
    fechas = set()
    for semilla in range(60):
        r = generar_formato_servilla("1", "20261229", [_fila("1", "F&S", "E")], random.Random(semilla))
        fechas.add(load_workbook(io.BytesIO(r.contenido)).active["F2"].value)
    assert fechas == {"20261231", "20270101", "20270102", "20270103", "20270104"}


def test_formato_servilla_se_puede_subir_al_generador():
    r = generar_formato_servilla("123791", "20260803", HISTO, random.Random(1))
    dat = generar_dat("123791", "20260731", [("servilla.xlsx", r.contenido)], HISTO, "terceros")
    lineas = dat.contenido_dat.decode("latin-1").split("\n")
    assert dat.registros == 4
    assert any(l.startswith("20260731DEV0520260803000001672696010") for l in lineas)


def test_orden_solo_prindel_lecta():
    with pytest.raises(ValueError, match="PRINDEL y LECTA"):
        generar_formato_servilla("1", "20260803", [_fila("1", "PRINDEL", "E")])
