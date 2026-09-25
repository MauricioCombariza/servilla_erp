import io
import zipfile

import pytest
from openpyxl import load_workbook

from app.services.generar_dat_service import construir_linea
from app.services.informe_global_service import (
    ItemInforme,
    asignar_dpto,
    generar_informe_global,
    parsear_centralizado,
    parsear_entregas,
)

import pandas as pd


def _reg(serial, estado, causal, courrier, dpto="ARAUCA"):
    return {
        "identdes": "8878", "oficina": "300", "nombred": "X", "dir_pred": "CL 1", "barrd1": "",
        "ciudad1": "ARAUCA", "dpto1": dpto, "Estado": estado, "Causal_Dev": causal,
        "F_recepcio": "20260808", "guias": serial, "F_GESTION": "20260810",
        "courrier": courrier, "orden": "123797",
    }


REGISTROS = [
    _reg("1672760018", "ENT", "00", "SERVILLA"),
    _reg("1672760019", "DEV", "05", "SERVILLA"),
    _reg("1672760020", "DEV", "07", "F&S"),
    _reg("7875000001", "ENT", "00", "LECTA"),
    _reg("7875000002", "DEV", "13", "LECTA"),
    _reg("1500000001", "ENT", "00", "PRINDEL"),
    _reg("1500000002", "DEV", "09", "PRINDEL"),
]
DAT_CENTRALIZADO = (
    "*BCSEXTCLP0220260804".ljust(312) + "\n"
    + "\n".join(construir_linea(r, "20260804") for r in REGISTROS) + "\n"
    + "*00000007".ljust(312) + "NOC\n"
).encode("latin-1")
DAT_ENTREGAS = (
    "0000001549DEV032026090100000167312813820260904\n"
    "0000009818ENT002026090100000167312926520260907\n"
    "0000010027DEV072026090100000167313135020260905\n"
).encode("latin-1")


def _libro(resultado, i=0):
    z = zipfile.ZipFile(io.BytesIO(resultado.contenido_zip))
    return load_workbook(io.BytesIO(z.read(resultado.items[i].nombre_excel)))


def test_parsea_centralizado_con_courier():
    df, detalle = parsear_centralizado(DAT_CENTRALIZADO)
    assert detalle == 7 and len(df) == 7
    assert list(df["courier"].unique()) == ["Servilla", "F&S", "Lecta", "Prindel"]
    assert df.iloc[0]["serial"] == "000001672760018"
    assert df.iloc[0]["fecha_corte"] == "20260804"


def test_parsea_entregas():
    df, detalle = parsear_entregas(DAT_ENTREGAS)
    assert detalle == 3 and len(df) == 3
    assert df.iloc[0][["estadoBCS", "fecha_inicio", "serial", "fecha_cierra"]].tolist() == [
        "03", "20260901", "000001673128138", "20260904"
    ]


def test_control_de_recepcion_en_filas_de_operador():
    item = ItemInforme("BCS_CLP_EXT_02_20260804.dat", DAT_CENTRALIZADO, "123797", "CLP_01", "centralizado")
    r = generar_informe_global([item], {})
    ws = _libro(r)["CONTROL DE RECEPCION"]

    assert ws["A1"].value == "Orden: 123797"
    assert ws["A2"].value == "Producto: CLP_01"
    assert ws["A3"].value == "Corte: 20260804"
    assert ws["C3"].value == "Fecha mínima: 2026-08-08"
    assert [c.value for c in ws[9][1:6]] == [None] * 5  # DOMINA sin registros
    assert [c.value for c in ws[10][1:6]] == [2, 1, 1, 0, 0]  # LECTA
    assert [c.value for c in ws[11][1:6]] == [3, 1, 1, 1, 0]  # SERVILLA + aliado F&S
    assert [c.value for c in ws[12][1:6]] == [2, 1, 0, 0, 1]  # PRINDEL
    assert ws["F13"].value == "=SUM(F9:F12)"
    # nada anexado debajo de la tabla
    assert all(c.value is None for fila in ws.iter_rows(min_row=14) for c in fila)


def test_causales_en_fila_del_producto():
    item = ItemInforme("a.dat", DAT_CENTRALIZADO, "123797", "clp 01", "centralizado")
    r = generar_informe_global([item], {"123797": [{"serial": "1672760018", "dpto1": "ARAUCA"}]})
    ws = _libro(r)["CAUSALES DE DEVOLUCION"]

    assert ws["A16"].value == "CLP_01"
    # códigos 5, 7, 9 y 13 → columnas G, I, K y L
    assert [ws[f"{c}16"].value for c in "CDEFGHIJKL"] == [None, None, None, None, 1, None, 1, None, 1, 1]
    assert ws["M16"].value == "=SUM(C16:L16)-I16-K16"
    # Los valores de ejemplo de la plantilla no quedan en otras filas
    assert all(ws[f"C{f}"].value is None for f in range(14, 23))
    assert r.items[0].advertencias == []


def test_producto_nuevo_agrega_fila_y_entregas_va_a_servilla():
    item = ItemInforme("bcs_TCR_EXT_02_20260901.dat", DAT_ENTREGAS, "123854", "TC_02", "entregas")
    r = generar_informe_global([item], {})
    wb = _libro(r)
    control, causales = wb["CONTROL DE RECEPCION"], wb["CAUSALES DE DEVOLUCION"]

    assert control["A3"].value == "Corte: 20260901"
    assert [c.value for c in control[11][1:6]] == [3, 1, 1, 1, 0]
    assert causales["A23"].value == "TC_02"
    assert causales["E23"].value == 1 and causales["I23"].value == 1
    assert any("se agregó una fila nueva" in a for a in r.items[0].advertencias)


def test_zip_con_un_libro_por_dat_y_consolidado():
    items = [
        ItemInforme("a.dat", DAT_CENTRALIZADO, "123797", "CLP_01", "centralizado"),
        ItemInforme("b.dat", DAT_CENTRALIZADO, "123797", "CLP_01", "centralizado"),
    ]
    histo = {"123797": [{"serial": "1672760018", "dpto1": "ARAUCA"}]}
    r = generar_informe_global(items, histo)
    nombres = zipfile.ZipFile(io.BytesIO(r.contenido_zip)).namelist()
    assert nombres == ["informe_CLP_01_123797.xlsx", "informe_CLP_01_123797_2.xlsx", "consolidado_departamentos.xlsx"]
    dptos = {(d["courier"], d["dpto1"]): d["cantidad"] for d in r.items[0].departamentos}
    assert dptos[("Servilla", "ARAUCA")] == 1
    assert dptos[("F&S", "Dev_Inicial")] == 1
    assert dptos[("Prindel", "Dev_Inicial")] == 1


def test_asignar_dpto_por_sufijo():
    histo = [{"serial": "001672760018", "dpto1": "META"}, {"serial": "99123", "dpto1": "CESAR"}]
    s = asignar_dpto(pd.Series(["1672760018", "123", "555"]), histo)
    assert s.tolist() == ["META", "CESAR", "SANTANDER"]


def test_dat_sin_registros_validos():
    with pytest.raises(ValueError, match="no se encontraron registros"):
        generar_informe_global([ItemInforme("x.dat", b"*solo cabecera\n", "1", "X", "centralizado")], {})
