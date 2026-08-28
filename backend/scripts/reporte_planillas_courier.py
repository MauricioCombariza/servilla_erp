"""
Reporte Excel de todas las planillas de un courier/transportadora (cod_mensajero).

Junta, para cada planilla histórica del código dado:
  - cantidades (local/nacional/total)
  - fecha de la planilla
  - valores (local/nacional/total)
  - número de factura (si ya fue facturada)
  - estado: Pagado / Se debe / Vencida / No facturado

Fuentes (PostgreSQL, servilla_erp):
  - seriales_gestion + personal        -> descubre TODAS las planillas del código
  - prefactura_planillas + prefacturas_courier + facturas_courier_cxp
        -> planillas que ya fueron agrupadas/facturadas
  - seriales_gestion (agregado)        -> planillas que aún no se agruparon en ninguna
        prefactura (mismo cálculo que _PLANILLAS_BASE_SQL en pagos_ciudades.py)

Uso:
    python reporte_planillas_courier.py                 # cod 1006, sale a Downloads
    python reporte_planillas_courier.py --cod 1006
    python reporte_planillas_courier.py --cod 1006 --output /ruta/personalizada.xlsx

Variables de entorno (túnel SSH ya abierto hacia el VPS):
    PostgreSQL — tunnel: ssh -L 5440:127.0.0.1:5440 root@204.168.150.196
        PG_HOST         (default: 127.0.0.1)
        PG_PORT         (default: 5440)
        PG_USER         (default: servilla)
        PG_PASSWORD     (requerida)
        PG_DB           (default: servilla_erp)
"""

import argparse
import os
import sys
from datetime import date

import psycopg2
import psycopg2.extras
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

PG_HOST = os.environ.get("PG_HOST", "127.0.0.1")
PG_PORT = int(os.environ.get("PG_PORT", 5440))
PG_USER = os.environ.get("PG_USER", "servilla")
PG_PASS = os.environ.get("PG_PASSWORD", "")
PG_DB = os.environ.get("PG_DB", "servilla_erp")

DOWNLOADS_DIR = "/mnt/c/Users/mcomb/Downloads"

HEADERS = [
    "Planilla", "Fecha planilla", "Cantidad local", "Cantidad nacional", "Cantidad total",
    "Valor local", "Valor nacional", "Valor total",
    "N° Factura", "Fecha factura", "Fecha vencimiento", "Estado", "Fecha de pago",
]

_TODAS_LAS_PLANILLAS_SQL = """
    SELECT DISTINCT sg.planilla
    FROM seriales_gestion sg
    JOIN personal p ON p.codigo = sg.cod_men
    WHERE sg.cod_men = %s
      AND p.tipo_personal IN ('courier_externo', 'transportadora')
      AND sg.estado != 'anulado'
"""

_PLANILLAS_INCLUIDAS_SQL = """
    SELECT
        pp.planilla, pp.fecha_escaner, pp.cantidad_local, pp.cantidad_nacional,
        pp.valor_local, pp.valor_nac, pp.valor_total,
        pf.estado AS prefactura_estado,
        cxp.numero_factura, cxp.fecha_emision, cxp.fecha_vencimiento,
        cxp.estado AS cxp_estado, cxp.fecha_pago
    FROM prefactura_planillas pp
    JOIN prefacturas_courier pf ON pf.id = pp.prefactura_id
    LEFT JOIN facturas_courier_cxp cxp ON cxp.prefactura_id = pf.id
    WHERE pf.cod_mensajero = %s
"""

# Agregación de planillas que aún no se agruparon en ninguna prefactura
# (misma lógica que _PLANILLAS_BASE_SQL en app/routers/pagos_ciudades.py, sin filtro de fecha).
_PLANILLAS_SIN_PREFACTURA_SQL = """
    SELECT
        sg.planilla,
        MIN(sg.f_esc) AS fecha_escaner,
        COUNT(*) FILTER (WHERE sg.ambito = 'bogota')   AS cantidad_local,
        COUNT(*) FILTER (WHERE sg.ambito = 'nacional') AS cantidad_nacional,
        COALESCE(SUM(sg.precio_mensajero) FILTER (WHERE sg.ambito = 'bogota'), 0)   AS valor_local,
        COALESCE(SUM(sg.precio_mensajero) FILTER (WHERE sg.ambito = 'nacional'), 0) AS valor_nac,
        COALESCE(SUM(sg.precio_mensajero), 0) AS valor_total
    FROM seriales_gestion sg
    JOIN personal p ON p.codigo = sg.cod_men
    WHERE sg.cod_men = %s
      AND p.tipo_personal IN ('courier_externo', 'transportadora')
      AND sg.estado != 'anulado'
      AND sg.planilla = ANY(%s)
    GROUP BY sg.planilla
"""


def connect_pg():
    return psycopg2.connect(host=PG_HOST, port=PG_PORT, user=PG_USER, password=PG_PASS, dbname=PG_DB)


def estado_reporte(row):
    """Deriva el estado visible del reporte a partir del estado de CxP (si existe)."""
    cxp_estado = row.get("cxp_estado")
    if cxp_estado == "pagada":
        return "Pagado"
    if cxp_estado == "vencida":
        return "Vencida (se debe)"
    if cxp_estado == "pendiente":
        return "Se debe"
    return "No facturado"


def obtener_filas(cur, cod):
    cur.execute(_TODAS_LAS_PLANILLAS_SQL, (cod,))
    todas = [r["planilla"] for r in cur.fetchall()]

    cur.execute(_PLANILLAS_INCLUIDAS_SQL, (cod,))
    incluidas = cur.fetchall()
    planillas_incluidas = {r["planilla"] for r in incluidas}

    faltantes = [p for p in todas if p not in planillas_incluidas]
    sin_prefactura = []
    if faltantes:
        cur.execute(_PLANILLAS_SIN_PREFACTURA_SQL, (cod, faltantes))
        sin_prefactura = cur.fetchall()

    filas = []
    for r in incluidas:
        filas.append(dict(r))
    for r in sin_prefactura:
        d = dict(r)
        d["prefactura_estado"] = None
        d["numero_factura"] = None
        d["fecha_emision"] = None
        d["fecha_vencimiento"] = None
        d["cxp_estado"] = None
        d["fecha_pago"] = None
        filas.append(d)

    filas.sort(key=lambda r: (r["fecha_escaner"] is None, r["fecha_escaner"]))
    return filas


def escribir_excel(cod, nombre_mensajero, filas, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Planillas {cod}"

    titulo = f"Planillas courier {cod}" + (f" - {nombre_mensajero}" if nombre_mensajero else "")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=13)

    header_row = 3
    for col, titulo_col in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=col, value=titulo_col)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    money_cols = {6, 7, 8}  # Valor local, Valor nacional, Valor total
    row_idx = header_row + 1
    for r in filas:
        cantidad_local = r["cantidad_local"] or 0
        cantidad_nacional = r["cantidad_nacional"] or 0
        ws.cell(row=row_idx, column=1, value=r["planilla"])
        ws.cell(row=row_idx, column=2, value=r["fecha_escaner"])
        ws.cell(row=row_idx, column=3, value=cantidad_local)
        ws.cell(row=row_idx, column=4, value=cantidad_nacional)
        ws.cell(row=row_idx, column=5, value=cantidad_local + cantidad_nacional)
        ws.cell(row=row_idx, column=6, value=float(r["valor_local"] or 0))
        ws.cell(row=row_idx, column=7, value=float(r["valor_nac"] or 0))
        ws.cell(row=row_idx, column=8, value=float(r["valor_total"] or 0))
        ws.cell(row=row_idx, column=9, value=r.get("numero_factura"))
        ws.cell(row=row_idx, column=10, value=r.get("fecha_emision"))
        ws.cell(row=row_idx, column=11, value=r.get("fecha_vencimiento"))
        ws.cell(row=row_idx, column=12, value=estado_reporte(r))
        ws.cell(row=row_idx, column=13, value=r.get("fecha_pago"))
        for col in money_cols:
            ws.cell(row=row_idx, column=col).number_format = "#,##0"
        for col in (2, 10, 11, 13):
            ws.cell(row=row_idx, column=col).number_format = "yyyy-mm-dd"
        row_idx += 1

    # Fila de totales
    total_row = row_idx
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=sum(r["cantidad_local"] or 0 for r in filas)).font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=sum(r["cantidad_nacional"] or 0 for r in filas)).font = Font(bold=True)
    ws.cell(
        row=total_row, column=5,
        value=sum((r["cantidad_local"] or 0) + (r["cantidad_nacional"] or 0) for r in filas),
    ).font = Font(bold=True)
    for col, key in ((6, "valor_local"), (7, "valor_nac"), (8, "valor_total")):
        c = ws.cell(row=total_row, column=col, value=sum(float(r[key] or 0) for r in filas))
        c.font = Font(bold=True)
        c.number_format = "#,##0"

    widths = [14, 14, 12, 14, 12, 14, 14, 14, 16, 14, 16, 18, 14]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cod", default="1006", help="Código de mensajero/courier (default: 1006)")
    parser.add_argument("--output", default=None, help="Ruta del Excel de salida")
    args = parser.parse_args()

    if not PG_PASS:
        sys.exit("Falta credencial: exporta PG_PASSWORD antes de correr este script.")

    output_path = args.output or os.path.join(
        DOWNLOADS_DIR, f"planillas_courier_{args.cod}_{date.today().isoformat()}.xlsx"
    )

    conn = connect_pg()
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cur.execute(
            "SELECT nombre_completo, tipo_personal FROM personal WHERE codigo = %s",
            (args.cod,),
        )
        personal = cur.fetchone()
        if personal is None:
            sys.exit(f"No existe personal con codigo={args.cod!r}.")
        print(f"Courier: {args.cod} - {personal['nombre_completo']} ({personal['tipo_personal']})")

        filas = obtener_filas(cur, args.cod)
        print(f"Planillas encontradas: {len(filas)}")

        no_facturadas = sum(1 for r in filas if estado_reporte(r) == "No facturado")
        se_debe = sum(1 for r in filas if estado_reporte(r) in ("Se debe", "Vencida (se debe)"))
        pagadas = sum(1 for r in filas if estado_reporte(r) == "Pagado")
        print(f"  No facturadas: {no_facturadas} | Se debe: {se_debe} | Pagadas: {pagadas}")

        escribir_excel(args.cod, personal["nombre_completo"], filas, output_path)
        print(f"\nExcel generado en: {output_path}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
