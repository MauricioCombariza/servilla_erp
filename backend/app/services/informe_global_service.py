"""
Informe global BCS: port web de infGlobal.sh (informeBCS.py + informeSimpleBCS.jl +
informeComplejoBCS.jl + consolidar_informes.py).

Por cada .dat (centralizado o entregas) llena la plantilla de BCS:
- CONTROL DE RECEPCION: encabezado (orden, producto, corte, fecha mínima) y las
  métricas por operador en sus filas (DOMINA, LECTA, SERVILLA, PRINDEL).
- CAUSALES DE DEVOLUCION: conteo por causal en la fila del producto.

Además arma un consolidado por departamento (dpto1 de bases_web.histo) de todos
los .dat. Los couriers aliados (Envigoex, F&S, J.R.G., POSTAL_COL...) cuentan
dentro de SERVILLA en la plantilla.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.services.excel_utils import construir_excel

RUTA_PLANTILLA = Path(__file__).resolve().parent.parent / "assets" / "plantilla_informe_bcs.xlsx"

TIPO_CENTRALIZADO = "centralizado"
TIPO_ENTREGAS = "entregas"
TIPOS = (TIPO_CENTRALIZADO, TIPO_ENTREGAS)

OPERADORES = ("DOMINA", "LECTA", "SERVILLA", "PRINDEL")
OPERADOR_DEFAULT = "SERVILLA"
HOJA_CONTROL = "CONTROL DE RECEPCION"
HOJA_CAUSALES = "CAUSALES DE DEVOLUCION"
ESTADO_ENTREGA = 0
ESTADO_DEV_INICIAL = 7
ESTADO_NRD = 9
DPTO_DEFAULT = "SANTANDER"
DPTO_DEV_INICIAL = "Dev_Inicial"
LONGITUD_MIN_SERIAL = 10
FORMATO_NUMERO = "#,##0"

# Centralizado: bloque de gestión (fecha_corte, estado, estadoBCS, fecha_inicio,
# serial, fecha_cierra) ubicado por la primera fecha AAAAMMDD válida de la línea,
# seguido opcionalmente de 'NOC' + '1<COURIER>'. El formato terceros (44
# caracteres) es solo el bloque, sin NOC ni courier.
ESTRUCTURA_CENTRALIZADO = [8, 3, 2, 8, 15, 8]
COLUMNAS_CENTRALIZADO = ["fecha_corte", "estado", "estadoBCS", "fecha_inicio", "serial", "fecha_cierra"]
PATRON_FECHA = re.compile(r"\b20\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])")
# Entregas: ancho fijo desde el inicio de la línea.
ESTRUCTURA_ENTREGAS = [10, 3, 2, 8, 15, 8]
COLUMNAS_ENTREGAS = ["serial_courier", "estado", "estadoBCS", "fecha_inicio", "serial", "fecha_cierra"]


@dataclass
class ItemInforme:
    nombre_archivo: str
    contenido: bytes
    orden: str
    nombre: str
    tipo: str


@dataclass
class ResultadoItem:
    nombre_archivo: str
    orden: str
    nombre: str
    tipo: str
    registros: int
    corte: str
    fecha_minima: str
    operadores: list[dict]
    causales: dict[str, int]
    nombre_excel: str
    contenido_excel: bytes
    advertencias: list[str] = field(default_factory=list)
    departamentos: list[dict] = field(default_factory=list)


@dataclass
class ResultadoGlobal:
    items: list[ResultadoItem]
    nombre_zip: str
    contenido_zip: bytes


def _cortar(texto: str, estructura: list[int]) -> list[str]:
    partes, inicio = [], 0
    for ancho in estructura:
        partes.append(texto[inicio:inicio + ancho].strip())
        inicio += ancho
    return partes


def _normalizar_courier(nombre: str) -> str:
    """SERVILLA -> Servilla; nombres con símbolos (J.R.G., F&S, POSTAL_COL) quedan tal cual."""
    return nombre.capitalize() if nombre.isalpha() else nombre


def _lineas(contenido: bytes) -> list[str]:
    return contenido.decode("latin-1").splitlines()


def parsear_centralizado(contenido: bytes) -> tuple[pd.DataFrame, int]:
    """Devuelve los registros y cuántas líneas de detalle (sin '*') tenía el archivo."""
    ancho = sum(ESTRUCTURA_CENTRALIZADO)
    filas, detalle = [], 0
    for linea in _lineas(contenido):
        if not linea.strip() or linea.startswith("*"):
            continue
        detalle += 1
        m = PATRON_FECHA.search(linea)
        if not m or len(linea) < m.start() + ancho:
            continue
        inicio = m.start()
        partes = _cortar(linea[inicio:inicio + ancho], ESTRUCTURA_CENTRALIZADO)
        resto = linea[inicio + ancho:]
        courier = ""
        if resto.startswith("NOC"):
            m_courier = re.match(r"\d?(\S+)", resto[3:])
            courier = _normalizar_courier(m_courier.group(1)) if m_courier else ""
        filas.append([*partes, courier])
    df = pd.DataFrame(filas, columns=[*COLUMNAS_CENTRALIZADO, "courier"])
    return df, detalle


def parsear_entregas(contenido: bytes) -> tuple[pd.DataFrame, int]:
    ancho_min = sum(ESTRUCTURA_ENTREGAS[:-1])
    filas, detalle = [], 0
    for linea in _lineas(contenido):
        if not linea.strip() or linea.startswith("*"):
            continue
        detalle += 1
        if len(linea) < ancho_min:
            continue
        filas.append([*_cortar(linea, ESTRUCTURA_ENTREGAS), ""])
    df = pd.DataFrame(filas, columns=[*COLUMNAS_ENTREGAS, "courier"])
    return df, detalle


def operador_de(courier: str) -> str:
    c = courier.strip().upper()
    return c if c in OPERADORES else OPERADOR_DEFAULT


def calcular_operadores(df: pd.DataFrame) -> list[dict]:
    estados = pd.to_numeric(df["estadoBCS"], errors="coerce").fillna(-1).astype(int)
    operadores = df["courier"].map(operador_de)
    resultado = []
    for op in OPERADORES:
        e = estados[operadores == op]
        resultado.append({
            "operador": op,
            "enviada": int(len(e)),
            "entrega": int((e == ESTADO_ENTREGA).sum()),
            "devoluciones": int((~e.isin([ESTADO_ENTREGA, ESTADO_DEV_INICIAL, ESTADO_NRD, -1])).sum()),
            "dev_iniciales": int((e == ESTADO_DEV_INICIAL).sum()),
            "nrd": int((e == ESTADO_NRD).sum()),
        })
    return resultado


def calcular_causales(df: pd.DataFrame) -> dict[int, int]:
    estados = pd.to_numeric(df["estadoBCS"], errors="coerce").dropna().astype(int)
    conteo = estados[estados != ESTADO_ENTREGA].value_counts().sort_index()
    return {int(k): int(v) for k, v in conteo.items()}


def _fecha_texto(valor: str) -> str:
    try:
        return datetime.strptime(valor, "%Y%m%d").strftime("%Y-%m-%d")
    except (TypeError, ValueError):
        return "N/A"


def _clave_producto(nombre) -> str:
    texto = unicodedata.normalize("NFKD", str(nombre or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def _llenar_control(ws, orden: str, nombre: str, corte: str, fecha_minima: str, operadores: list[dict]):
    ws["A1"] = f"Orden: {orden}"
    ws["A2"] = f"Producto: {nombre}"
    ws["A3"] = f"Corte: {corte}"
    ws["C3"] = f"Fecha mínima: {fecha_minima}"

    fila_operador, fila_total = {}, None
    for fila in range(1, ws.max_row + 1):
        etiqueta = str(ws.cell(row=fila, column=1).value or "").strip().upper()
        if etiqueta in OPERADORES:
            fila_operador[etiqueta] = fila
        elif etiqueta == "TOTAL":
            fila_total = fila
    faltan = [op for op in OPERADORES if op not in fila_operador]
    if faltan or fila_total is None:
        raise ValueError(f"La plantilla no tiene las filas de operador esperadas ({', '.join(faltan) or 'TOTAL'})")

    claves = ["enviada", "entrega", "devoluciones", "dev_iniciales", "nrd"]
    for op in operadores:
        fila = fila_operador[op["operador"]]
        for col, clave in enumerate(claves, start=2):
            celda = ws.cell(row=fila, column=col)
            # DOMINA sin registros queda en blanco, como en el informe manual
            celda.value = op[clave] if op["enviada"] else None
            celda.number_format = FORMATO_NUMERO

    primera, ultima = min(fila_operador.values()), max(fila_operador.values())
    for col in range(2, 2 + len(claves)):
        letra = ws.cell(row=fila_total, column=col).column_letter
        celda = ws.cell(row=fila_total, column=col)
        celda.value = f"=SUM({letra}{primera}:{letra}{ultima})"
        celda.number_format = FORMATO_NUMERO
    for fila in (4, 6):
        ws.cell(row=fila, column=2).number_format = FORMATO_NUMERO


def _llenar_causales(ws, nombre: str, causales: dict[int, int]) -> list[str]:
    advertencias = []
    fila_header = next(
        (f for f in range(1, ws.max_row + 1)
         if str(ws.cell(row=f, column=1).value or "").strip().lower() == "producto"),
        None,
    )
    if fila_header is None:
        raise ValueError(f"La plantilla no tiene el encabezado 'Producto' en la hoja {HOJA_CAUSALES}")

    col_codigo, col_total = {}, None
    for col in range(1, ws.max_column + 1):
        valor = ws.cell(row=fila_header, column=col).value
        if isinstance(valor, (int, float)):
            col_codigo[int(valor)] = col
        elif str(valor or "").strip().upper() == "TOTAL":
            col_total = col

    # La plantilla trae valores de ejemplo: se limpian todas las filas de producto
    filas_producto = [
        f for f in range(fila_header + 1, ws.max_row + 1) if ws.cell(row=f, column=1).value
    ]
    for f in filas_producto:
        for col in col_codigo.values():
            ws.cell(row=f, column=col).value = None

    clave = _clave_producto(nombre)
    fila = next((f for f in filas_producto if _clave_producto(ws.cell(row=f, column=1).value) == clave), None)
    if fila is None:
        fila = (filas_producto[-1] if filas_producto else fila_header) + 1
        molde = filas_producto[-1] if filas_producto else fila_header
        for col in range(1, ws.max_column + 1):
            origen = ws.cell(row=molde, column=col)
            if origen.has_style:
                ws.cell(row=fila, column=col)._style = copy(origen._style)
        ws.cell(row=fila, column=1).value = nombre
        ws.cell(row=fila, column=2).value = "BCS"
        advertencias.append(f"El producto '{nombre}' no estaba en la hoja de causales; se agregó una fila nueva")

    for codigo, cantidad in causales.items():
        col = col_codigo.get(codigo)
        if col is None:
            advertencias.append(f"Causal {codigo:02d} ({cantidad}) no tiene columna en la plantilla")
            continue
        ws.cell(row=fila, column=col).value = cantidad

    if col_total is not None and col_codigo:
        letra = lambda c: ws.cell(row=fila, column=c).column_letter  # noqa: E731
        primera, ultima = letra(min(col_codigo.values())), letra(max(col_codigo.values()))
        formula = f"=SUM({primera}{fila}:{ultima}{fila})"
        # Devoluciones iniciales (07) y NRD (09) no cuentan como devolución
        for codigo in (ESTADO_DEV_INICIAL, ESTADO_NRD):
            if codigo in col_codigo:
                formula += f"-{letra(col_codigo[codigo])}{fila}"
        ws.cell(row=fila, column=col_total).value = formula
    return advertencias


def construir_excel_item(orden: str, nombre: str, corte: str, fecha_minima: str,
                         operadores: list[dict], causales: dict[int, int]) -> tuple[bytes, list[str]]:
    wb = load_workbook(RUTA_PLANTILLA)
    _llenar_control(wb[HOJA_CONTROL], orden, nombre, corte, fecha_minima, operadores)
    advertencias = _llenar_causales(wb[HOJA_CAUSALES], nombre, causales)
    wb.active = wb.sheetnames.index(HOJA_CONTROL)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), advertencias


def _mapa_dpto(filas_histo: list[dict]) -> dict[str, str]:
    mapa = {}
    for f in filas_histo:
        serial = str(f.get("serial") or "").strip().lstrip("0")
        if serial and serial not in mapa:
            mapa[serial] = str(f.get("dpto1") or "").strip() or DPTO_DEFAULT
    return mapa


def asignar_dpto(seriales: pd.Series, filas_histo: list[dict]) -> pd.Series:
    """dpto1 por serial: coincidencia exacta y, si no, un serial de histo que termine
    en el del .dat (como el cruce por cadenas invertidas de informeComplejoBCS.jl)."""
    mapa = _mapa_dpto(filas_histo)
    por_sufijo: dict[int, dict[str, str]] = {}

    def buscar(serial: str) -> str:
        if serial in mapa:
            return mapa[serial]
        n = len(serial)
        if n == 0:
            return DPTO_DEFAULT
        if n not in por_sufijo:
            por_sufijo[n] = {}
            for s, d in mapa.items():
                if len(s) > n:
                    por_sufijo[n].setdefault(s[-n:], d)
        return por_sufijo[n].get(serial, DPTO_DEFAULT)

    return seriales.map(buscar)


def calcular_departamentos(df: pd.DataFrame, tipo: str, filas_histo: list[dict]) -> list[dict]:
    d = df.copy()
    d["serial_norm"] = d["serial"].str.strip().str.lstrip("0")
    if tipo == TIPO_CENTRALIZADO:
        d = d[d["serial_norm"].str.len() >= LONGITUD_MIN_SERIAL].drop_duplicates("serial_norm")
    if d.empty:
        return []
    d["dpto1"] = asignar_dpto(d["serial_norm"], filas_histo)
    estados = pd.to_numeric(d["estadoBCS"], errors="coerce").fillna(0).astype(int)
    d.loc[estados.isin([ESTADO_DEV_INICIAL, ESTADO_NRD]), "dpto1"] = DPTO_DEV_INICIAL
    d["courier"] = d["courier"].where(d["courier"] != "", OPERADOR_DEFAULT.capitalize())
    conteo = d.groupby(["courier", "dpto1"]).size().reset_index(name="cantidad")
    return conteo.sort_values(["courier", "dpto1"]).to_dict("records")


def _nombre_seguro(texto: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", texto).strip("_") or "informe"


def procesar_item(item: ItemInforme, filas_histo: list[dict]) -> ResultadoItem:
    if item.tipo not in TIPOS:
        raise ValueError(f"{item.nombre_archivo}: tipo inválido ({item.tipo})")
    parsear = parsear_centralizado if item.tipo == TIPO_CENTRALIZADO else parsear_entregas
    df, detalle = parsear(item.contenido)
    if df.empty:
        raise ValueError(f"{item.nombre_archivo}: no se encontraron registros válidos para el tipo {item.tipo}")

    advertencias = []
    if len(df) != detalle:
        advertencias.append(
            f"El archivo tiene {detalle} líneas de detalle pero se procesaron {len(df)} registros"
        )

    # Centralizado: el corte es fecha_corte y la fecha mínima la menor fecha_inicio
    # (F_recepcio). Entregas no trae esas columnas: el corte es fecha_inicio y la
    # fecha mínima la menor fecha de cierre.
    if item.tipo == TIPO_CENTRALIZADO:
        col_corte, col_minima = "fecha_corte", "fecha_inicio"
    else:
        col_corte, col_minima = "fecha_inicio", "fecha_cierra"
    cortes = sorted(c for c in df[col_corte].unique() if c)
    corte = cortes[0] if cortes else "N/A"
    if len(cortes) > 1:
        advertencias.append(f"El archivo tiene varias fechas de corte: {', '.join(cortes)}")
    fechas = pd.to_datetime(df[col_minima], format="%Y%m%d", errors="coerce").dropna()
    fecha_minima = fechas.min().strftime("%Y-%m-%d") if not fechas.empty else "N/A"

    operadores = calcular_operadores(df)
    causales = calcular_causales(df)
    contenido_excel, adv_excel = construir_excel_item(
        item.orden, item.nombre, corte, fecha_minima, operadores, causales
    )
    advertencias += adv_excel
    if not filas_histo:
        advertencias.append(f"La orden {item.orden} no tiene registros en histo: departamentos en {DPTO_DEFAULT}")

    return ResultadoItem(
        nombre_archivo=item.nombre_archivo,
        orden=item.orden,
        nombre=item.nombre,
        tipo=item.tipo,
        registros=len(df),
        corte=corte,
        fecha_minima=fecha_minima,
        operadores=operadores,
        causales={f"{k:02d}": v for k, v in causales.items()},
        nombre_excel=f"informe_{_nombre_seguro(item.nombre)}_{item.orden}.xlsx",
        contenido_excel=contenido_excel,
        advertencias=advertencias,
        departamentos=calcular_departamentos(df, item.tipo, filas_histo),
    )


def construir_consolidado(resultados: list[ResultadoItem]) -> bytes:
    filas = [
        {"Orden": r.orden, "Proceso": r.nombre, "Tipo": r.tipo, "Archivo": r.nombre_archivo,
         "Courier": d["courier"], "Departamento": d["dpto1"], "Cantidad": d["cantidad"]}
        for r in resultados for d in r.departamentos
    ]
    columnas = ["Orden", "Proceso", "Tipo", "Archivo", "Courier", "Departamento", "Cantidad"]
    return construir_excel("Consolidado por departamento", columnas, filas, [10, 16, 14, 36, 14, 26, 10])


def generar_informe_global(items: list[ItemInforme], histo_por_orden: dict[str, list[dict]]) -> ResultadoGlobal:
    if not items:
        raise ValueError("Agrega al menos un archivo .dat")
    resultados = [procesar_item(it, histo_por_orden.get(it.orden, [])) for it in items]

    usados: set[str] = set()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for r in resultados:
            nombre, n = r.nombre_excel, 2
            while nombre in usados:
                nombre = r.nombre_excel.replace(".xlsx", f"_{n}.xlsx")
                n += 1
            usados.add(nombre)
            r.nombre_excel = nombre
            zf.writestr(nombre, r.contenido_excel)
        zf.writestr("consolidado_departamentos.xlsx", construir_consolidado(resultados))

    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return ResultadoGlobal(items=resultados, nombre_zip=f"informe_global_{fecha}.zip", contenido_zip=buffer.getvalue())
