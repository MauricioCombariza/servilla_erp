from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas.devoluciones import DevolucionDocumentoItem

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets" / "devoluciones"
HEADER_LOGO = ASSETS_DIR / "header_logo.png"
FOOTER_GRAPHIC = ASSETS_DIR / "footer_graphic.png"

MESES_ES = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]


def _fecha_es(d: date) -> str:
    return f"{d.day} de {MESES_ES[d.month - 1]} de {d.year}"


def construir_docx_devolucion(
    items: list[DevolucionDocumentoItem],
    generado_por: str | None = None,
) -> bytes:
    doc = Document()

    section = doc.sections[0]
    if HEADER_LOGO.exists():
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_p.add_run().add_picture(str(HEADER_LOGO), width=Inches(5.5))
    if FOOTER_GRAPHIC.exists():
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_p.add_run().add_picture(str(FOOTER_GRAPHIC), width=Inches(1.2))

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("ACTA DE DEVOLUCIÓN")
    run.bold = True
    run.font.size = Pt(18)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo.add_run(f"Fecha: {_fecha_es(date.today())}")

    if generado_por:
        generado = doc.add_paragraph()
        generado.alignment = WD_ALIGN_PARAGRAPH.CENTER
        generado.add_run(f"Generado por: {generado_por}").italic = True

    doc.add_paragraph()

    columnas = ["Serial", "Nombre", "Dirección", "Localidad"]
    tabla = doc.add_table(rows=1, cols=len(columnas))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = tabla.rows[0].cells
    for cell, nombre_col in zip(header_cells, columnas):
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(nombre_col)
        r.bold = True

    for item in items:
        cells = tabla.add_row().cells
        cells[0].text = item.serial
        cells[1].text = item.nombre or ""
        cells[2].text = item.direccion or ""
        cells[3].text = item.localidad or ""

    doc.add_paragraph()
    doc.add_paragraph()

    firma = doc.add_paragraph()
    firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    firma.add_run("_" * 40)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
