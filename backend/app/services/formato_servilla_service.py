"""
Formato Excel de gestión para Servilla: prellena, a partir de la orden en
bases_web.histo, las columnas que luego lee el generador del .dat
(serial, Estado, Causal_Dev, F_recepcio, guias, F_GESTION) más `motivo`.

Solo se incluyen los courriers propios (se excluyen PRINDEL y LECTA, que envían
su propio Excel). La causal se deduce del motivo; los motivos no reconocidos
quedan sin causal para ajustarlos a mano.
"""
from __future__ import annotations

import io
import random
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.generar_dat_service import COLUMNAS_EXCEL

COURRIERS_EXCLUIDOS = {"PRINDEL", "LECTA"}
COLUMNAS_FORMATO = [*COLUMNAS_EXCEL, "motivo"]
DIAS_GESTION_MIN = 2
DIAS_GESTION_MAX = 6

# Causal BCS → nombres con que aparece el motivo en histo. Se comparan
# normalizados (sin tildes ni signos, en minúsculas) y admitiendo que histo
# trunca el motivo a 15 caracteres ("Direccion Errad", "Zona Alto Riesg").
CAUSALES: dict[str, list[str]] = {
    "00": ["entrega"],
    "01": ["desconocido", "destinatario n"],
    "02": ["rehusado", "rehusada"],
    "03": ["traslado"],
    "05": ["direccion errada", "no cubrimiento"],
    "06": ["direccion incompleta", "dir incompleta"],
    "07": ["inconsistencia", "inconsistente"],
    "08": ["zona alto riesgo"],
    "09": ["faltante", "sobrante", "docto en camino"],
    "13": ["cerrado"],
    "14": ["siniestro"],
}
_LARGO_MIN_TRUNCADO = 6

# Mapeo de cod_sec cuando el serial no tiene retorno/ret_esc (igual a updateHisto.py)
_MOTIVO_POR_COD_SEC = {
    "DEV_ 1": "Traslado",
    "DEV_ 2": "Direccion Errada",
    "DEV_ 5": "Dir Incompleta",
    "DEV_17": "Desconocido",
    "DEV_21": "Inconsistente",
    "DEV_22": "Rehusada",
    "DEV_33": "Dir Incompleta",
    "DEV_36": "Zona Alto Riesg",
    "DEV_48": "No Cubrimiento",
}


@dataclass
class ResultadoFormato:
    nombre: str
    contenido: bytes
    filas: int
    excluidos: int
    # [{serial, courrier, motivo}] de las filas que quedaron sin causal
    sin_causal: list[dict] = field(default_factory=list)


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def motivo_histo(fila: dict) -> str:
    """Motivo de gestión del serial, con la misma regla de updateHisto.py."""
    retorno = (fila.get("retorno") or "").strip()
    ret_esc = (fila.get("ret_esc") or "").strip()
    if "D" in (retorno, ret_esc):
        motivo = (fila.get("motivo") or "").strip()
    elif retorno == "o":
        motivo = "Dirección Errada"
    elif "E" in (retorno, ret_esc):
        motivo = "Entrega"
    else:
        motivo = ""
    if motivo in ("", "0"):
        cod_sec = (fila.get("cod_sec") or "").strip().lstrip("*")
        motivo = _MOTIVO_POR_COD_SEC.get(cod_sec, "")
    return motivo


def causal_de_motivo(motivo: str) -> str:
    """Causal BCS del motivo, o '' si no se reconoce."""
    valor = _normalizar(motivo)
    if not valor:
        return ""
    for causal, nombres in CAUSALES.items():
        for nombre in nombres:
            if valor == nombre:
                return causal
            if len(valor) >= _LARGO_MIN_TRUNCADO and nombre.startswith(valor):
                return causal
            if valor.startswith(nombre + " "):
                return causal
    return ""


def _sumar_dias(fecha: str, dias: int) -> str:
    return (datetime.strptime(fecha, "%Y%m%d") + timedelta(days=dias)).strftime("%Y%m%d")


def generar_formato_servilla(
    orden: str,
    f_recepcio: str,
    filas_histo: list[dict],
    rng: random.Random | None = None,
) -> ResultadoFormato:
    """f_recepcio en AAAAMMDD. F_GESTION = F_recepcio + 2..6 días calendario por fila."""
    if not filas_histo:
        raise ValueError(f"La orden {orden} no tiene registros en histo")
    rng = rng or random.Random()

    vistos: set[str] = set()
    filas: list[dict] = []
    excluidos = 0
    for fila in filas_histo:
        serial = str(fila.get("serial") or "").strip()
        courrier = (fila.get("courrier") or "").strip()
        if not serial or serial in vistos:
            continue
        vistos.add(serial)
        if courrier.upper() in COURRIERS_EXCLUIDOS:
            excluidos += 1
            continue
        motivo = motivo_histo(fila)
        causal = causal_de_motivo(motivo)
        filas.append({
            "serial": serial,
            "Estado": ("ENT" if causal == "00" else "DEV") if causal else "",
            "Causal_Dev": causal,
            "F_recepcio": f_recepcio,
            "guias": serial,
            "F_GESTION": _sumar_dias(f_recepcio, rng.randint(DIAS_GESTION_MIN, DIAS_GESTION_MAX)),
            "motivo": motivo,
            "courrier": courrier,
        })

    if not filas:
        raise ValueError(f"La orden {orden} no tiene seriales fuera de PRINDEL y LECTA")

    filas.sort(key=lambda f: f["serial"])
    sin_causal = [
        {"serial": f["serial"], "courrier": f["courrier"], "motivo": f["motivo"]}
        for f in filas if not f["Causal_Dev"]
    ]
    return ResultadoFormato(
        nombre=f"formato_servilla_{orden}_{f_recepcio}.xlsx",
        contenido=_construir_excel(filas),
        filas=len(filas),
        excluidos=excluidos,
        sin_causal=sin_causal,
    )


def _construir_excel(filas: list[dict]) -> bytes:
    """Encabezados en la fila 1 (como espera leer_excels), todo en texto y las
    filas sin causal resaltadas para completarlas a mano."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Gestion"
    resaltado = PatternFill("solid", fgColor="FFF2CC")

    for col, nombre in enumerate(COLUMNAS_FORMATO, start=1):
        ws.cell(row=1, column=col, value=nombre).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = 22 if nombre == "motivo" else 16

    for fila_idx, fila in enumerate(filas, start=2):
        for col, nombre in enumerate(COLUMNAS_FORMATO, start=1):
            c = ws.cell(row=fila_idx, column=col, value=fila[nombre])
            c.number_format = "@"
            if not fila["Causal_Dev"]:
                c.fill = resaltado
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
